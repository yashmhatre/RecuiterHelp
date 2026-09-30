"""Read a resume file and suggest form values.

Suggestions, not an import. The plan requires a human to check every profile, and a pre-filled
form the person corrects before saving is the cheapest way to make that the default rather than
a separate review step nobody does.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

__all__ = ["ParsedResume", "parse_resume", "extract_profile_hints"]

SKILL_VOCAB = (
    "python", "java", "javascript", "typescript", "golang", "rust", "c++", "c#", "php", "ruby",
    "scala", "kotlin", "swift", "django", "flask", "fastapi", "spring boot", "spring", "node.js",
    "nodejs", "react", "angular", "vue", "next.js", "express", "postgresql", "postgres", "mysql",
    "mongodb", "redis", "elasticsearch", "kafka", "rabbitmq", "celery", "docker", "kubernetes",
    "aws", "azure", "gcp", "terraform", "ansible", "jenkins", "github actions", "ci/cd",
    "graphql", "rest api", "microservices", "machine learning", "deep learning", "pytorch",
    "tensorflow", "scikit-learn", "nlp", "llm", "pandas", "numpy", "spark", "airflow", "dbt",
    "snowflake", "tableau", "power bi", "sql", "linux", "git", "selenium", "pytest", "cypress",
)


@dataclass
class ParsedResume:
    text: str
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)


def parse_resume(path: Path | str) -> ParsedResume:
    """Extract plain text from PDF, DOCX or TXT. Never raises on a bad file."""
    path = Path(path)
    suffix = path.suffix.lower()

    if not path.is_file():
        return ParsedResume("", warnings=["file_not_found"])
    if path.stat().st_size == 0:
        return ParsedResume("", warnings=["empty_file"])

    try:
        if suffix == ".pdf":
            return _parse_pdf(path)
        if suffix == ".docx":
            return _parse_docx(path)
        if suffix in {".txt", ".rtf"}:
            return ParsedResume(_normalise(path.read_text(encoding="utf-8", errors="replace")))
        if suffix == ".doc":
            return ParsedResume("", warnings=["legacy_doc_unsupported"])
    except Exception as exc:  # noqa: BLE001 - a bad upload must not 500 the app
        return ParsedResume("", warnings=[f"parse_failed: {type(exc).__name__}"])

    return ParsedResume("", warnings=[f"unsupported_suffix: {suffix}"])


def _parse_pdf(path: Path) -> ParsedResume:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except PdfReadError as exc:
        return ParsedResume("", warnings=[f"unreadable_pdf: {exc}"])

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            return ParsedResume("", warnings=["encrypted_pdf"])

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            pages.append("")

    text = _normalise("\n".join(pages))
    warnings = [] if len(text) > 40 else ["no_extractable_text"]
    return ParsedResume(text, page_count=len(reader.pages), warnings=warnings)


def _parse_docx(path: Path) -> ParsedResume:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    text = _normalise("\n".join(parts))
    return ParsedResume(text, warnings=[] if len(text) > 40 else ["no_extractable_text"])


def _normalise(text: str) -> str:
    """Rejoin words broken across line ends, collapse whitespace, keep line structure."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _clean_phone(match) -> str:
    """Only return a phone number long enough to plausibly be one."""
    if not match:
        return ""
    value = match.group(0).strip()
    return value if len(value) >= 8 else ""


def extract_profile_hints(text: str) -> dict:
    """Guess name, email, phone, title, skills, years and location from resume text."""
    lowered = text.lower()

    email_match = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone_match = re.search(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)?\d{5,10}", text)

    name = ""
    for line in text.splitlines()[:6]:
        stripped = line.strip()
        if (
            2 <= len(stripped.split()) <= 4
            and "@" not in stripped
            and not re.search(r"\d", stripped)
            and stripped.replace(" ", "").replace(".", "").replace("-", "").isalpha()
        ):
            name = " ".join(w.capitalize() for w in stripped.split())
            break

    title = ""
    title_match = re.search(
        r"\b((?:senior|lead|principal|staff|junior)?\s*(?:python|java|backend|front[\s-]?end|"
        r"full[\s-]?stack|devops|data|ml|machine learning|qa|cloud|software)\s*"
        r"(?:engineer|developer|architect|scientist|analyst|consultant))\b",
        lowered,
    )
    if title_match:
        title = " ".join(w.capitalize() for w in title_match.group(1).split())

    years = None
    for pattern in (
        r"(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:relevant\s+|total\s+)?experience",
        r"experience\s*[:\-]?\s*(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)",
        r"(\d+(?:\.\d+)?)\s*\+\s*(?:years?|yrs?)",
    ):
        found = re.search(pattern, lowered)
        if found:
            years = float(found.group(1))
            break

    location = ""
    for city in (
        "pune", "bangalore", "bengaluru", "hyderabad", "mumbai", "delhi", "noida", "gurgaon",
        "chennai", "kolkata", "ahmedabad", "jaipur", "indore", "remote",
    ):
        if re.search(rf"\b{city}\b", lowered):
            location = city.title()
            break

    skills = sorted({s for s in SKILL_VOCAB if re.search(rf"(?<![a-z]){re.escape(s)}(?![a-z])", lowered)})

    summary = ""
    summary_match = re.search(
        r"(?:summary|profile|objective|about)\s*[:\n]\s*(.{60,400}?)(?:\n\n|\nexperience|\nskills)",
        text,
        re.IGNORECASE | re.DOTALL,
    )
    if summary_match:
        summary = " ".join(summary_match.group(1).split())
    elif text:
        summary = " ".join(text.split()[:45])

    return {
        "name": name,
        "email": email_match.group(0) if email_match else "",
        "phone": _clean_phone(phone_match),
        "title": title,
        "skills": ", ".join(skills),
        "years_experience": years,
        "location": location,
        "summary": summary,
    }
